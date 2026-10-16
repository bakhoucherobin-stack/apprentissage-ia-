"""Importe le tableau des allergènes rempli par le restaurant dans js/data.js.

Usage : python3 outils/importer-allergenes.py chemin/vers/allergenes-pizza-food-a-remplir.xlsx
Seules les lignes marquées « OUI » dans la colonne « Vérifié » sont publiées ;
les autres produits affichent « appelez-nous avant de commander ».
(Nécessite openpyxl : pip install openpyxl)
"""
import json, re, sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / 'js' / 'data.js'
COLS = ['ARACHIDES', 'CELERI', 'CRUSTACES', 'FRUITS A COQUES', 'GLUTEN', 'LAIT', 'LUPIN', 'MOLLUSQUES',
        'MOUTARDE', 'OEUFS', 'POISSONS', 'SESAME', 'SOJA', 'SULFITES']

ws = load_workbook(sys.argv[1], data_only=True)['Allergènes']
rows = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r[1] or str(r[17] or '').strip().upper() != 'OUI':
        continue
    bits = ''.join('1' if str(v or '').strip().upper() == 'X' else '0' for v in r[3:17])
    rows.append([str(r[1]).strip().upper(), bits])

src = DATA.read_text(encoding='utf-8')
head, body = src.split('\n', 1)
data = json.loads(re.sub(r'^window\.PF_DATA = |;\s*$', '', body.strip()))
data['allergens'] = {'cols': COLS, 'rows': rows}
DATA.write_text(head + '\nwindow.PF_DATA = ' + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ';\n', encoding='utf-8')
print(f'{len(rows)} produits vérifiés importés dans js/data.js')
